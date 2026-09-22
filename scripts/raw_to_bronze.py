"""Load multi-format landing files into a raw Bronze SQLite database."""

from __future__ import annotations

import argparse
import csv
import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator

from docx import Document
from openpyxl import load_workbook


FIELDS = [
    "order_id", "order_date", "customer_id", "customer_name", "customer_email",
    "customer_phone", "product_id", "product_name", "category", "quantity",
    "unit_price", "discount_pct", "payment_method", "shipping_status", "province",
    "total_amount",
]
SUPPORTED_FILES = {
    "customer_orders.csv", "customer_orders.json", "customer_orders.txt",
    "customer_orders.xlsx", "customer_orders.docx",
}


def as_record(values: dict[str, Any]) -> dict[str, Any]:
    return {field: values.get(field) for field in FIELDS}


def read_csv(path: Path) -> Iterator[dict[str, Any]]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        yield from csv.DictReader(file)


def read_json(path: Path) -> Iterator[dict[str, Any]]:
    yield from json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> Iterator[dict[str, Any]]:
    with path.open(encoding="utf-8") as file:
        for line in file:
            if line.strip():
                yield json.loads(line)


def read_xlsx(path: Path) -> Iterator[dict[str, Any]]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    for values in rows:
        yield dict(zip(headers, values))


def read_docx(path: Path) -> Iterator[dict[str, Any]]:
    table = Document(path).tables[0]
    headers = [cell.text for cell in table.rows[0].cells]
    for row in table.rows[1:]:
        yield dict(zip(headers, (cell.text for cell in row.cells)))


def records_from(path: Path) -> Iterator[dict[str, Any]]:
    readers = {
        ".csv": read_csv,
        ".json": read_json,
        ".txt": read_jsonl,
        ".xlsx": read_xlsx,
        ".docx": read_docx,
    }
    yield from readers[path.suffix](path)


def create_table(connection: sqlite3.Connection) -> None:
    columns = ",\n        ".join(f'"{field}" TEXT' for field in FIELDS)
    connection.execute(
        f"""CREATE TABLE IF NOT EXISTS bronze_customer_orders (
        bronze_id INTEGER PRIMARY KEY AUTOINCREMENT,
        source_file TEXT NOT NULL,
        source_row INTEGER NOT NULL,
        extracted_at TEXT NOT NULL,
        raw_json TEXT NOT NULL,
        {columns}
        )"""
    )


def load_landing(input_dir: Path, database: Path) -> None:
    database.parent.mkdir(parents=True, exist_ok=True)
    files = [path for path in sorted(input_dir.iterdir()) if path.name in SUPPORTED_FILES]
    if not files:
        raise FileNotFoundError(f"No supported landing files found in {input_dir}")

    extracted_at = datetime.now(timezone.utc).isoformat()
    with sqlite3.connect(database) as connection:
        create_table(connection)
        connection.execute("DELETE FROM bronze_customer_orders")
        columns = ", ".join(["source_file", "source_row", "extracted_at", "raw_json", *FIELDS])
        placeholders = ", ".join("?" for _ in range(4 + len(FIELDS)))
        for path in files:
            for source_row, values in enumerate(records_from(path), start=1):
                record = as_record(values)
                connection.execute(
                    f"INSERT INTO bronze_customer_orders ({columns}) VALUES ({placeholders})",
                    [path.name, source_row, extracted_at, json.dumps(values, ensure_ascii=False),
                     *[record[field] for field in FIELDS]],
                )
            print(f"loaded {path.name}")
        total = connection.execute("SELECT COUNT(*) FROM bronze_customer_orders").fetchone()[0]
    print(f"bronze_customer_orders: {total:,} rows -> {database}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=Path("data/landing/customer_orders"))
    parser.add_argument("--database", type=Path, default=Path("data/warehouse/warehouse.db"))
    args = parser.parse_args()
    load_landing(args.input_dir, args.database)


if __name__ == "__main__":
    main()