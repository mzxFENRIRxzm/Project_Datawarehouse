"""Adapt Downloads/project's Northwind model to MinIO + Iceberg + Trino.

Each input content hash gets immutable Silver/Gold tables. Only the final
dashboard view is switched after validation; retries never append duplicate rows.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

FILES = ('customers.csv', 'products.json', 'orders_2026.xlsx',
         'order_details.txt', 'shippers_policy.docx')
MODEL_VERSION = 'v1'


def decimal(value):
    result = Decimal(str(value))
    if not result.is_finite():
        raise ValueError('Non-finite numeric value')
    return result


def unique(rows, key):
    seen = {}
    for row in rows:
        identity = row[key]
        if identity in (None, ''):
            raise ValueError(f'Missing {key}')
        if identity in seen and seen[identity] != row:
            raise ValueError(f'Conflicting duplicate {key}: {identity}')
        seen[identity] = row
    return list(seen.values())


def validate(data):
    for table, key in [('customers', 'customer_id'), ('products', 'product_id'),
                       ('orders', 'order_id'), ('shippers', 'shipper_id')]:
        if not data[table]:
            raise ValueError(f'Empty source: {table}')
        data[table] = unique(data[table], key)
    keys = {name: {r[key] for r in data[name]} for name, key in
            [('customers', 'customer_id'), ('products', 'product_id'),
             ('orders', 'order_id'), ('shippers', 'shipper_id')]}
    for row in data['orders']:
        if row['customer_id'] not in keys['customers'] or row['shipper_id'] not in keys['shippers']:
            raise ValueError(f'Orphan order: {row["order_id"]}')
        if row['freight'] < 0:
            raise ValueError('Negative freight')
    if not data['order_details']:
        raise ValueError('Empty order details')
    for row in data['order_details']:
        if row['order_id'] not in keys['orders'] or row['product_id'] not in keys['products']:
            raise ValueError(f'Orphan order detail: {row["line_id"]}')
        if row['quantity'] <= 0 or row['unit_price'] < 0 or not 0 <= row['discount'] <= 1:
            raise ValueError(f'Invalid sales measure: {row["line_id"]}')


def read_sources(root):
    from docx import Document
    from openpyxl import load_workbook
    for name in FILES:
        if not (root / name).is_file():
            raise FileNotFoundError(root / name)
    with (root / 'customers.csv').open(encoding='utf-8-sig', newline='') as f:
        # Contact names and phone numbers stay in the original Bronze files only.
        customers = [dict(customer_id=r['CustomerID'].strip(),
                          company_name=r['CompanyName'].strip(), city=r['City'].strip())
                     for r in csv.DictReader(f)]
    products = [dict(product_id=int(r['ProductID']), product_name=r['ProductName'].strip(),
                     category_id=int(r['CategoryID']), unit_price=decimal(r['Specs']['UnitPrice']))
                for r in json.loads((root / 'products.json').read_text(encoding='utf-8'))]
    workbook = load_workbook(root / 'orders_2026.xlsx', read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows)
        orders = []
        for values in rows:
            r = dict(zip(headers, values))
            raw_date = r['OrderDate']
            day = raw_date.date() if isinstance(raw_date, datetime) else date.fromisoformat(str(raw_date))
            orders.append(dict(order_id=int(r['OrderID']), customer_id=str(r['CustomerID']).strip(),
                               order_date=day, shipper_id=int(r['ShipperID']), freight=decimal(r['Freight'])))
    finally:
        workbook.close()
    with (root / 'order_details.txt').open(encoding='utf-8-sig', newline='') as f:
        details = [dict(line_id=i, order_id=int(r['OrderID']), product_id=int(r['ProductID']),
                        unit_price=decimal(r['UnitPrice']), quantity=int(r['Quantity']),
                        discount=decimal(r['Discount']))
                   for i, r in enumerate(csv.DictReader(f, delimiter='|'), start=1)]
    shippers = []
    for paragraph in Document(root / 'shippers_policy.docx').paragraphs:
        match = re.search(r'SHIP_ID:\s*(\d+)\s*->\s*(.+)', paragraph.text)
        if match:
            shippers.append(dict(shipper_id=int(match[1]), shipper_name=match[2].strip()))
    data = dict(customers=customers, products=products, orders=orders,
                order_details=details, shippers=shippers)
    validate(data)
    return data


def source_hash(root):
    digest = hashlib.sha256(MODEL_VERSION.encode())
    for name in FILES:
        digest.update(name.encode())
        digest.update((root / name).read_bytes())
    return digest.hexdigest()[:16]


def archive(root, batch):
    import boto3
    client = boto3.client('s3', endpoint_url=os.environ['MINIO_ENDPOINT'],
                          aws_access_key_id=os.environ['LAKEHOUSE_MINIO_USER'],
                          aws_secret_access_key=os.environ['LAKEHOUSE_MINIO_PASSWORD'],
                          region_name='us-east-1')
    for name in FILES:
        client.upload_file(str(root / name), 'warehouse', f'bronze/northwind/{batch}/{name}')


def stage(data, schema):
    import psycopg2
    from psycopg2 import sql
    from psycopg2.extras import execute_values
    types = {'customer_id': 'text', 'company_name': 'text', 'city': 'text',
             'product_id': 'bigint', 'product_name': 'text', 'category_id': 'bigint',
             'unit_price': 'numeric(18,2)', 'order_id': 'bigint', 'order_date': 'date',
             'shipper_id': 'bigint', 'shipper_name': 'text', 'freight': 'numeric(18,2)',
             'line_id': 'bigint', 'quantity': 'bigint', 'discount': 'numeric(6,4)'}
    connection = psycopg2.connect(host=os.environ['POSTGRES_HOST'], dbname=os.environ['POSTGRES_DB'],
                                 user=os.environ['POSTGRES_USER'], password=os.environ['POSTGRES_PASSWORD'])
    try:
        with connection, connection.cursor() as cursor:
            cursor.execute(sql.SQL('CREATE SCHEMA IF NOT EXISTS {}').format(sql.Identifier(schema)))
            for table, rows in data.items():
                target = sql.Identifier(schema, table)
                columns = list(rows[0])
                definitions = sql.SQL(', ').join(sql.SQL('{} {}').format(sql.Identifier(c), sql.SQL(types[c])) for c in columns)
                cursor.execute(sql.SQL('CREATE TABLE IF NOT EXISTS {} ({})').format(target, definitions))
                cursor.execute(sql.SQL('SELECT COUNT(*) FROM {}').format(target))
                count = cursor.fetchone()[0]
                if count == 0:
                    execute_values(cursor, sql.SQL('INSERT INTO {} VALUES %s').format(target),
                                   [[r[c] for c in columns] for r in rows], page_size=1000)
                elif count != len(rows):
                    raise ValueError(f'Staging row count mismatch: {table}')
    finally:
        connection.close()


def publish(data, batch):
    from trino.dbapi import connect
    connection = connect(host=os.environ.get('TRINO_HOST', 'localhost'),
                         port=int(os.environ.get('TRINO_PORT', '8180')), user='northwind_etl')
    def query(statement):
        cursor = connection.cursor()
        try:
            cursor.execute(statement)
            return cursor.fetchall()
        finally:
            cursor.close()
    silver = f'northwind_silver_{batch}'
    gold = f'northwind_gold_{batch}'
    try:
        for schema in (silver, gold, 'northwind_gold'):
            query(f"CREATE SCHEMA IF NOT EXISTS iceberg.{schema} WITH (location = 's3://warehouse/{schema}/')")
        for table, rows in data.items():
            query(f"CREATE TABLE IF NOT EXISTS iceberg.{silver}.{table} WITH (format='PARQUET') AS SELECT * FROM postgresql.northwind_stage_{batch}.{table}")
            if query(f'SELECT count(*) FROM iceberg.{silver}.{table}')[0][0] != len(rows):
                raise ValueError(f'Silver count mismatch: {table}')
        models = {
            'dim_customer': f'SELECT * FROM iceberg.{silver}.customers',
            'dim_product': f'SELECT * FROM iceberg.{silver}.products',
            'dim_shipper': f'SELECT * FROM iceberg.{silver}.shippers',
            'dim_date': f'''SELECT DISTINCT order_date, year(order_date) AS year,
                month(order_date) AS month, week(order_date) AS iso_week,
                year_of_week(order_date) AS iso_year FROM iceberg.{silver}.orders''',
            # Freight belongs to an order, never repeat it on every sales line.
            'fact_orders': f'''SELECT *, CAST(round(freight * DECIMAL '1.07', 2) AS DECIMAL(18,2))
                AS freight_including_vat FROM iceberg.{silver}.orders''',
            'fact_sales': f'''SELECT d.*, o.customer_id, o.order_date, o.shipper_id,
                CAST(round(d.quantity * d.unit_price * (1 - d.discount), 2) AS DECIMAL(18,2)) AS net_revenue
                FROM iceberg.{silver}.order_details d JOIN iceberg.{silver}.orders o ON d.order_id=o.order_id''',
        }
        for name, statement in models.items():
            query(f"CREATE TABLE IF NOT EXISTS iceberg.{gold}.{name} WITH (format='PARQUET') AS {statement}")
        count, revenue = query(f'SELECT count(*), sum(net_revenue) FROM iceberg.{gold}.fact_sales')[0]
        expected = sum((r['quantity'] * r['unit_price'] * (1-r['discount'])).quantize(Decimal('.01'), rounding='ROUND_HALF_UP')
                       for r in data['order_details'])
        if count != len(data['order_details']) or revenue != expected:
            raise ValueError(f'Gold reconciliation failed: {count}, {revenue}, expected {expected}')
        # One atomic view replacement publishes the complete validated batch for BI.
        query(f'''CREATE OR REPLACE VIEW iceberg.northwind_gold.sales_dashboard AS
            SELECT f.*, p.product_name, p.category_id, c.city, s.shipper_name
            FROM iceberg.{gold}.fact_sales f
            JOIN iceberg.{gold}.dim_product p ON f.product_id=p.product_id
            JOIN iceberg.{gold}.dim_customer c ON f.customer_id=c.customer_id
            JOIN iceberg.{gold}.dim_shipper s ON f.shipper_id=s.shipper_id''')
        result = dict(batch=batch, source_counts={k: len(v) for k, v in data.items()},
                      sales_lines=count, net_revenue=str(revenue), gold_schema=gold,
                      dashboard='iceberg.northwind_gold.sales_dashboard')
        print(json.dumps(result, indent=2))
        return result
    finally:
        connection.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-dir', type=Path, default=Path('data/landing/northwind_thai_large_data'))
    parser.add_argument('--validate-only', action='store_true')
    args = parser.parse_args()
    data = read_sources(args.input_dir)
    batch = source_hash(args.input_dir)
    print(f'Validated batch {batch}: ' + ', '.join(f'{k}={len(v)}' for k, v in data.items()))
    if not args.validate_only:
        archive(args.input_dir, batch)
        stage(data, f'northwind_stage_{batch}')
        publish(data, batch)


if __name__ == '__main__':
    main()
